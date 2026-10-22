# -*- coding: utf-8 -*-
"""Word 导出（python-docx）。

坑点：只设 font.name 中文仍会回落宋体，必须同时写 w:eastAsia。
"""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services import grading
from services.export import model as export_model

CN_FONT = "微软雅黑"
EN_FONT = "Microsoft YaHei"


def _set_run_font(run, size=None, bold=None, color=None):
    run.font.name = EN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _bottom_border(paragraph, sz=8, space=1):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.makeelement(qn("w:pBdr"), {})
    bottom = ppr.makeelement(
        qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): str(sz), qn("w:space"): str(space), qn("w:color"): "000000"}
    )
    pbdr.append(bottom)
    ppr.append(pbdr)


def _setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    return doc


def _para(doc, text="", size=11, bold=False, align=None, space_before=0, space_after=0, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = indent
    r = p.add_run(text)
    _set_run_font(r, size=size, bold=bold)
    return p


def _zero_cell_margins(table) -> None:
    """表格左右内边距清零，卷头三栏才能真的贴边对齐。"""
    tbl_pr = table._tbl.tblPr
    mar = tbl_pr.makeelement(qn("w:tblCellMar"), {})
    for side in ("left", "right"):
        mar.append(tbl_pr.makeelement(qn("w:" + side), {qn("w:w"): "0", qn("w:type"): "dxa"}))
    tbl_pr.append(mar)


def _header_row(doc, left: str, mid: str, right: str) -> None:
    """卷头一行三栏：考生信息靠左、题量与满分居中、得分靠右。"""
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    _zero_cell_margins(table)
    widths = (Cm(6.6), Cm(4.4), Cm(6.0))
    aligns = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)
    for i, text in enumerate((left, mid, right)):
        table.columns[i].width = widths[i]
        cell = table.cell(0, i)
        cell.width = widths[i]
        p = cell.paragraphs[0]
        p.text = ""
        p.alignment = aligns[i]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        _set_run_font(run, size=10.5)


def _options_table(doc, item: dict):
    opts = item["options"]
    if not opts:
        return
    two = len(opts) <= 6 and max(len(o["text"]) for o in opts) <= 18
    cols = 2 if two else 1
    rows = (len(opts) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    for idx, o in enumerate(opts):
        r, c = idx % rows, idx // rows
        cell = table.cell(r, c)
        cell.paragraphs[0].text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("%s. %s" % (o["label"], o["text"]))
        _set_run_font(run, size=11)
    # 表格后留一点间距
    _para(doc, "", size=6)


def build_docx(vm: dict, out_path: str) -> str:
    doc = _setup_document()

    _para(doc, vm["title"], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    mid = ["共 %d 题" % vm["total"]]
    if vm.get("show_score") and vm.get("score_total"):
        mid.append("满分 %s 分" % grading.fmt_score(vm["score_total"]))
    left = "考生信息：%s" % (vm.get("candidate") or "____________") if vm.get("show_candidate") else ""
    right = "得分：__________" if vm.get("show_score") else ""
    _header_row(doc, left, "　".join(mid), right)

    # 分隔线紧贴卷头那一行：极窄空段落 + 下边框，段前 0、边框间距 0
    sep = _para(doc, "", size=1, space_before=0, space_after=8)
    sep.paragraph_format.line_spacing = Pt(1)
    _bottom_border(sep, space=0)

    for g in vm["groups"]:
        _para(
            doc,
            "%s（%s）" % (g["name"], export_model.group_subtitle(g)),
            size=11.5,
            bold=True,
            space_before=10,
            space_after=5,
        )
        for it in g["items"]:
            # 题干必须带答题括号：单选/判断 2 字符间隙、多选 6 字符间隙
            stem = grading.ensure_answer_parens(it["stem"], it["qtype"])
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r1 = p.add_run("%d. " % it["pos"])
            _set_run_font(r1, size=11, bold=True)
            r2 = p.add_run(stem)
            _set_run_font(r2, size=11)
            _options_table(doc, it)

    if vm["with_answer"]:
        doc.add_page_break()
        _para(doc, "参考答案", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        line: list[str] = []
        for i, a in enumerate(vm["answer_sheet"], start=1):
            line.append("%d. %s" % (a["pos"], a["answer_text"]))
            if i % 6 == 0:
                _para(doc, "　　".join(line), size=11, space_after=4)
                line = []
        if line:
            _para(doc, "　　".join(line), size=11, space_after=4)
        _para(
            doc,
            "说明：多选题须与标准答案完全一致才计分（漏选、错选均为 0 分）。",
            size=9,
            space_before=10,
        )
        for run in doc.paragraphs[-1].runs:
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    _para(doc, "本试卷由「随题」生成", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16)
    for run in doc.paragraphs[-1].runs:
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.save(out_path)
    return out_path
