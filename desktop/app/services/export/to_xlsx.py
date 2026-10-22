# -*- coding: utf-8 -*-
"""Excel 导出（openpyxl）：试题 + 参考答案 两个工作表。"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services import grading

FONT = "微软雅黑"
HEAD_FILL = PatternFill("solid", fgColor="1F5FD0")
HEAD_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=13, bold=True, color="0F172A")
BODY_FONT = Font(name=FONT, size=10, color="1F2937")
THIN = Side(style="thin", color="C9D6E8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(ws, row: int, ncol: int):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 24


def build_xlsx(vm: dict, out_path: str) -> str:
    max_opts = 0
    for g in vm["groups"]:
        for it in g["items"]:
            max_opts = max(max_opts, it["option_count"])
    max_opts = max(2, min(26, max_opts))

    wb = Workbook()
    ws = wb.active
    ws.title = "试题"

    opt_letters = [chr(65 + i) for i in range(max_opts)]
    headers = ["序号", "题型", "分值", "题干"] + ["选项" + x for x in opt_letters] + ["考生作答"]
    ncol = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    tcell = ws.cell(row=1, column=1, value=vm["title"])
    tcell.font = TITLE_FONT
    tcell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 卷头一行三栏：考生信息靠左、题量与满分居中、得分靠右
    mid_bits = ["共 %d 题" % vm["total"]]
    if vm.get("show_score") and vm.get("score_total"):
        mid_bits.append("满分 %s 分" % grading.fmt_score(vm["score_total"]))
    cand_text = ("考生信息：%s" % (vm.get("candidate") or "____________")) if vm.get("show_candidate") else ""
    score_text = "得分：__________" if vm.get("show_score") else ""

    sub_font = Font(name=FONT, size=9.5, color="475569")
    ws.row_dimensions[2].height = 20

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
    left_cell = ws.cell(row=2, column=1, value=cand_text)
    left_cell.font = sub_font
    left_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=ncol - 1)
    mid_cell = ws.cell(row=2, column=3, value="　".join(mid_bits))
    mid_cell.font = sub_font
    mid_cell.alignment = Alignment(horizontal="center", vertical="center")

    right_cell = ws.cell(row=2, column=ncol, value=score_text)
    right_cell.font = sub_font
    right_cell.alignment = Alignment(horizontal="right", vertical="center")

    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h)
    _style_header(ws, 3, ncol)

    row = 4
    for g in vm["groups"]:
        for it in g["items"]:
            ws.cell(row=row, column=1, value=it["pos"])
            ws.cell(row=row, column=2, value=g["type_label"])
            ws.cell(row=row, column=3, value=(it.get("score") or None))
            ws.cell(row=row, column=4, value=it["stem"])
            for k, o in enumerate(it["options"]):
                ws.cell(row=row, column=5 + k, value=o["text"])
            for c in range(1, ncol + 1):
                cell = ws.cell(row=row, column=c)
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="center" if c in (1, 2, 3, ncol) else "left",
                )
            ws.row_dimensions[row].height = 38
            row += 1

    widths = {1: 6, 2: 9, 3: 6, 4: 62}
    for i in range(len(opt_letters)):
        widths[5 + i] = 24
    widths[ncol] = 16
    for idx, w in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "E4"

    # ---- 参考答案 ----
    wa = wb.create_sheet("参考答案")
    a_headers = ["序号", "题型", "正确答案", "题干"]
    wa.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    acell = wa.cell(row=1, column=1, value="参考答案 · " + vm["title"])
    acell.font = TITLE_FONT
    acell.alignment = Alignment(horizontal="center", vertical="center")
    wa.row_dimensions[1].height = 28

    for i, h in enumerate(a_headers, start=1):
        wa.cell(row=2, column=i, value=h)
    _style_header(wa, 2, 4)

    r = 3
    type_map = {it["pos"]: g["type_label"] for g in vm["groups"] for it in g["items"]}
    stem_map = {it["pos"]: it["stem"] for g in vm["groups"] for it in g["items"]}
    for a in vm["answer_sheet"]:
        wa.cell(row=r, column=1, value=a["pos"])
        wa.cell(row=r, column=2, value=type_map.get(a["pos"], ""))
        wa.cell(row=r, column=3, value=a["answer_text"])
        wa.cell(row=r, column=4, value=stem_map.get(a["pos"], "")[:60])
        for c in range(1, 5):
            cell = wa.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top", wrap_text=True, horizontal="center" if c in (1, 2) else "left"
            )
        ans_cell = wa.cell(row=r, column=3)
        ans_cell.font = Font(name=FONT, size=10, bold=True, color="B42318")
        r += 1

    for idx, w in {1: 6, 2: 9, 3: 14, 4: 70}.items():
        wa.column_dimensions[get_column_letter(idx)].width = w
    wa.freeze_panes = "A3"

    wb.save(out_path)
    return out_path
