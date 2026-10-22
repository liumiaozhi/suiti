# -*- coding: utf-8 -*-
"""后端端到端自检（无界面）。用法：
    .venv-build\\Scripts\\python.exe tests\\run_check.py
覆盖：模板生成 → 三种写法导入 → 抽题打乱 → 练习判题 → 考试评卷 → 记录统计 → 三格式导出。
"""
import json
import os
import sys
import tempfile
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.join(os.path.dirname(HERE), "app")
sys.path.insert(0, APP)

SANDBOX = os.path.join(tempfile.gettempdir(), "suiti_check_%d" % os.getpid())
os.makedirs(SANDBOX, exist_ok=True)
os.environ["APPDATA"] = SANDBOX  # 数据目录指向沙箱，避免污染真实数据

PASS, FAIL = [], []


def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(("  [OK]   " if cond else "  [FAIL] ") + name + (("  " + str(extra)) if extra else ""))


def section(t):
    print("\n=== %s ===" % t)


def make_xlsx(path, rows, use_opt_cols=False):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"
    if use_opt_cols:
        head = ["题型", "题干", "选项A", "选项B", "选项C", "选项D", "正确答案"]
    else:
        head = ["题型", "题干", "选项", "正确答案"]
    ws.append(head)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def main():
    import db

    db.init_db()

    import api as api_mod
    from services import bank_import, session as sess_mod, stats

    A = api_mod.Api()

    # ---------------------------------------------------------------- 模板
    section("模板生成")
    tpl = os.path.join(SANDBOX, "tpl.xlsx")
    bank_import.build_template(tpl)
    check("模板文件生成", os.path.getsize(tpl) > 5000, "%d bytes" % os.path.getsize(tpl))

    from openpyxl import load_workbook

    wb = load_workbook(tpl)
    ws = wb["题库"]
    heads = [c.value for c in ws[1]]
    check("模板表头含 题型/题干/选项/正确答案", all(h in heads for h in ["题型", "题干", "选项", "正确答案"]), heads[:4])
    check("模板含填写说明页", "填写说明" in wb.sheetnames)

    # ---------------------------------------------------------------- 导入
    section("题库导入")
    f1 = make_xlsx(
        os.path.join(SANDBOX, "银行业务基础.xlsx"),
        [
            ["单选", "我国法定存款准备金制度由（　）负责统一实施。", "中国人民银行|金融监管总局|财政部|商业银行", "A"],
            ["单选", "商业银行资本充足率监管要求中，核心一级资本充足率不得低于（　）。", "4%|5%|6%|8%", "B"],
            ["单选", "下列属于银行表外业务的是（　）。", "贷款|同业存放|银行承兑汇票|买入返售", "C"],
            ["多选题", "下列属于流动性资产的有（　）。", "库存现金|存放央行款项|固定资产|一个月内到期拆放", "ABD"],
            ["判断题", "同业存单属于商业银行负债业务。", "", "√"],
        ],
    )
    r1 = A.import_bank_path({"path": f1})
    check("写法一（选项单列）导入成功", r1["ok"] and r1["data"]["ok"], r1.get("data", {}).get("errors"))
    if r1["ok"]:
        check("题量统计正确", r1["data"]["counts"] == {"total": 5, "single": 3, "multi": 1, "judge": 1}, r1["data"]["counts"])
        bank1 = r1["data"]["bank_id"]

    f2 = make_xlsx(
        os.path.join(SANDBOX, "合规与风控.xlsx"),
        [
            ["单选", "（A~H 分列写法）洗钱风险自评估的英文缩写是（　）。", "RBA", "SRE", "IRA", "AML", "A"],
            ["多选", "大额交易报告的情形包括（　）。", "现金收付5万元以上", "转账200万元以上", "跨境交易", "内部调拨", "ABC"],
            ["判断", "反洗钱内部控制制度可以两年修订一次。", "", "", "", "", "错误"],
        ],
        use_opt_cols=True,
    )
    r2 = A.import_bank_path({"path": f2})
    check("写法二（选项A~H 分列）导入成功", r2["ok"] and r2["data"]["ok"], r2.get("data", {}).get("errors"))
    if r2["ok"]:
        bank2 = r2["data"]["bank_id"]

    f3 = make_xlsx(
        os.path.join(SANDBOX, "判断题混合写法.xlsx"),
        [
            ["判断", "存贷比是流动性风险监管指标。（T/F 写法）", "", "T"],
            ["判断", "同业拆借期限可以超过一年。（× 写法）", "", "×"],
            ["判断", "债券投资属于资产业务。（对 写法）", "", "对"],
            ["判断", "信用风险只存在于贷款业务中。（错 写法）", "", "错"],
            ["判断", "净值型理财产品不承诺保本。（正确 写法）", "", "正确"],
        ],
    )
    r3 = A.import_bank_path({"path": f3})
    check("判断题五种答案写法全部识别", r3["ok"] and r3["data"]["ok"], r3.get("data", {}).get("errors"))

    f4 = make_xlsx(
        os.path.join(SANDBOX, "坏文件.xlsx"),
        [["单选", "答案超出选项范围（应被拒绝）", "甲|乙|丙|丁", "Z"]],
    )
    r4 = A.import_bank_path({"path": f4})
    check("答案越界整文件拒绝导入", r4["ok"] and (r4["data"]["ok"] is False), r4["data"]["errors"] if r4["ok"] else r4)
    banks_after = A.list_banks({})["data"]["banks"]
    check("坏文件未入库（题库数仍为 3）", len(banks_after) == 3, [b["name"] for b in banks_after])

    # ---------------------------------------------------------------- 抽题
    section("出题与打乱")
    settings = {
        "bank_id": bank1,
        "types": ["single", "multi", "judge"],
        "amount_mode": "all",
        "amount": {},
        "score": {"single": 2, "multi": 4, "judge": 1},
        "q_order": "asc",
        "a_order": "shuffle",
    }
    pv = A.preview_paper(settings)
    check(
        "预览：5 题、按题型分值合计满分 11",
        pv["ok"] and pv["data"]["total"] == 5 and pv["data"]["score_total"] == 11,
        {"total": pv["data"]["total"], "score_total": pv["data"]["score_total"], "by_type": pv["data"]["by_type"]},
    )
    check(
        "预览：各题型分值分别为 2 / 4 / 1",
        [pv["data"]["by_type"][t]["score"] for t in ("single", "multi", "judge")] == [2, 4, 1],
        pv["data"]["by_type"],
    )

    # 反复建 10 次卷，验证答案映射在随机打乱下始终自洽
    mismatch = 0
    for _ in range(10):
        cs = A.create_session(dict(settings, mode="practice"))
        sid = cs["data"]["session_id"]
        items = cs["data"]["items"]
        with db.ro() as con:
            rows = con.execute(
                "SELECT pos, answer_disp_json, answer_disp_text, option_order_json, options_json FROM session_items"
                " WHERE session_id=? ORDER BY pos",
                (sid,),
            ).fetchall()
        for r in rows:
            order = json.loads(r["option_order_json"])
            ans_disp = json.loads(r["answer_disp_json"])
            # 独立复算：展示序里答案位置对应的原始选项，必须等于原始正确答案那几项
            # 做法：用 answer_disp 反查 -> 展示选项；再与原标题比对由 answer_disp_text 校验
            item = next(i for i in items if i["pos"] == r["pos"])
            if item["qtype"] != "judge":
                letters = ",".join(chr(65 + k) for k in ans_disp)
                if letters != r["answer_disp_text"]:
                    mismatch += 1
        A.abandon_session({"session_id": sid})
    check("10 次随机打乱后答案展示文本一致", mismatch == 0, "mismatch=%d" % mismatch)

    # 严格校验一次映射正确性
    cs = A.create_session(dict(settings, mode="practice"))
    sid = cs["data"]["session_id"]
    with db.ro() as con:
        rows = con.execute("SELECT * FROM session_items WHERE session_id=? ORDER BY pos", (sid,)).fetchall()
    ok_map = True
    for r in rows:
        order = json.loads(r["option_order_json"])
        raw_opts = json.loads(r["options_json"])
        disp_opts = [raw_opts[i] for i in order]
        ans_disp = json.loads(r["answer_disp_json"])
        with db.ro() as con2:
            q = con2.execute("SELECT answer_json FROM questions WHERE id=?", (r["qid"],)).fetchone()
        raw_ans = json.loads(q["answer_json"])
        # 展示序下答案对应的选项文本，必须等于原始答案对应的选项文本
        if sorted(disp_opts[a] for a in ans_disp) != sorted(raw_opts[a] for a in raw_ans):
            ok_map = False
    check("打乱后答案指向的选项文本与题库一致", ok_map)

    # ---------------------------------------------------------------- 练习判题
    section("练习模式判题")
    pos_ans = {r["pos"]: json.loads(r["answer_disp_json"]) for r in rows}
    # 先故意把第 2 题答错，其余全对
    target_wrong = 2
    n_opts2 = len(json.loads(rows[1]["options_json"]))
    bad = [x for x in range(n_opts2) if x not in pos_ans[target_wrong]][:1]
    res_bad = A.answer({"session_id": sid, "pos": target_wrong, "chosen": bad})
    check(
        "答错立即返回错误并给出正确答案",
        res_bad["ok"] and res_bad["data"]["is_correct"] is False and bool(res_bad["data"]["answer_text"]),
        res_bad.get("data"),
    )
    for p, a in pos_ans.items():
        if p == target_wrong:
            continue
        res = A.answer({"session_id": sid, "pos": p, "chosen": a})
        if p == 1:
            check("答对立即返回正确", res["ok"] and res["data"]["is_correct"] is True, res.get("data"))
    dup = A.answer({"session_id": sid, "pos": 1, "chosen": [0]})
    check("练习模式同一题不能重复作答", dup["ok"] is False and dup["error"]["code"] == "BAD_INPUT", dup.get("error"))

    sub = A.submit_session({"session_id": sid})
    d = sub["data"]
    check("练习交卷：总题数 5", d["total"] == 5, d["total"])
    check("练习交卷：正确 4 / 错误 1 / 漏答 0", (d["right"], d["wrong"], d["missing"]) == (4, 1, 0), (d["right"], d["wrong"], d["missing"]))
    check("练习交卷：得分 9（错的那题是单选 2 分）", d["score"] == 9, d["score"])
    check("练习交卷：满分 11（2+2+2+4+1）", d["score_total"] == 11, d["score_total"])
    check("练习交卷：正确率 80%", d["accuracy"] == 80.0, d["accuracy"])
    check("练习交卷：带回各题型分值", d["scores"] == {"single": 2, "multi": 4, "judge": 1}, d.get("scores"))

    # ---------------------------------------------------------------- 考试评卷
    section("模拟/正式考试评卷")
    cs2 = A.create_session(dict(settings, mode="exam", candidate="张三 001", a_order="keep"))
    sid2 = cs2["data"]["session_id"]
    with db.ro() as con:
        rows2 = con.execute("SELECT * FROM session_items WHERE session_id=? ORDER BY pos", (sid2,)).fetchall()
    ans2 = {r["pos"]: json.loads(r["answer_disp_json"]) for r in rows2}
    # 答 3 题（全对），留 2 题不答
    for p in list(ans2)[:3]:
        A.answer({"session_id": sid2, "pos": p, "chosen": ans2[p]})
    st = A.get_session_state({"session_id": sid2})
    check("考试中已答 3 题", st["data"]["done_cnt"] == 3, st["data"]["done_cnt"])
    mk = A.toggle_mark({"session_id": sid2, "pos": 5, "marked": True})
    check("标记成功", mk["ok"] and mk["data"]["marked"] is True)
    sub2 = A.submit_session({"session_id": sid2})["data"]
    check("考试交卷：正确 3 / 漏答 2", (sub2["right"], sub2["wrong"], sub2["missing"]) == (3, 0, 2), (sub2["right"], sub2["wrong"], sub2["missing"]))
    check("考试交卷：得分 6 / 正确率 60%", (sub2["score"], sub2["accuracy"]) == (6, 60.0), (sub2["score"], sub2["accuracy"]))
    check("考试交卷：满分 11", sub2["score_total"] == 11, sub2["score_total"])
    res = A.get_result({"session_id": sid2})["data"]
    check("错题列表含 2 道漏答且标记为未作答", len(res["wrong_items"]) == 2 and all(w["missing"] for w in res["wrong_items"]), len(res["wrong_items"]))
    check("错题项带 qid（供归集重难错题）", all(w.get("qid") for w in res["wrong_items"]), [w.get("qid") for w in res["wrong_items"]])

    # 模拟考试不应进入考试记录
    cs3 = A.create_session(dict(settings, mode="mock", candidate="李四"))
    sid3 = cs3["data"]["session_id"]
    A.submit_session({"session_id": sid3})

    # ---------------------------------------------------------------- 记录
    section("记录统计")
    rec = A.get_records({})["data"]
    ov = rec["overall"]
    print("     做题记录整体:", ov)
    check("做题记录：题库分行数 >= 1", len(rec["by_bank"]) >= 1, len(rec["by_bank"]))
    check("做题记录：正确数 >= 7", ov["right"] >= 7, ov["right"])
    ex = A.list_exam_records({})["data"]
    print("     考试记录整体:", ex["overall"])
    check("考试记录只含正式考试（1 条）", ex["overall"]["exam_cnt"] == 1, ex["overall"]["exam_cnt"])
    check("考试记录题量合计 5", ex["overall"]["total_q"] == 5, ex["overall"]["total_q"])
    check("考试记录行含考生信息", ex["rows"] and ex["rows"][0]["candidate"] == "张三 001", ex["rows"][0] if ex["rows"] else None)

    # ---------------------------------------------------------------- 导出
    section("三格式导出")
    from services.export import exporter

    for fmt in ("docx", "xlsx", "pdf"):
        try:
            out = exporter.export_paper(sid2, fmt, window=None, with_answer=True)
            path = out.get("path")
            size = os.path.getsize(path) if path and os.path.exists(path) else 0
            check("导出 %s 成功" % fmt, bool(path) and size > 3000, "%s %d bytes" % (path, size))
            if fmt == "pdf" and path and os.path.exists(path):
                with open(path, "rb") as f:
                    check("PDF 魔数正确", f.read(4) == b"%PDF")
        except Exception as e:
            check("导出 %s 成功" % fmt, False, "%s: %s" % (type(e).__name__, e))
            traceback.print_exc()

    # 试卷里不应出现日期
    from datetime import datetime as _dt

    from services.export import model as export_model, html_tpl

    vm = export_model.build(sid2, True)
    html = html_tpl.render_html(vm)
    today = _dt.now()
    date_hits = [s for s in (today.strftime("%Y-%m-%d"), today.strftime("%Y/%m/%d"), today.strftime("%Y年%m月%d日")) if s in html]
    check("导出的试卷不含日期", not date_hits, date_hits)
    check("导出的试卷含各题型满分说明", "满分 11 分" in html, "含「满分 11 分」" if "满分 11 分" in html else "缺失")

    # 卷头规范：标题只做题库名、四项固定（题量/满分/考生信息/得分）
    check("卷面标题就是题库名", vm["title"] == "银行业务基础", vm["title"])
    check(
        "标题不含练习/模拟/正式等后缀",
        not any(w in vm["title"] for w in ("练习", "模拟", "正式考试", "重难错题")),
        vm["title"],
    )
    check("卷头有题量", "共 5 题" in html)
    check("卷头有考生信息（未填写也是空栏）", "考生信息：" in html)
    check("卷头有得分栏", "得分：" in html)

    # 练习卷也要保留考生信息
    vm_practice = export_model.build(sid, True)  # sid 是练习会话，无考生信息
    html_practice = html_tpl.render_html(vm_practice)
    check("练习卷标题也是题库名", vm_practice["title"] == "银行业务基础", vm_practice["title"])
    check("练习卷保留考生信息栏", "考生信息：" in html_practice and "line" in html_practice)
    check("练习卷也有得分栏", "得分：" in html_practice)

    # 答题括号：单选/判断留 2 字符，多选留 6 字符；没有就补在末尾
    from services import grading as _g

    check(
        "单选括号留 2 字符",
        _g.ensure_answer_parens("下列说法正确的是（　）", "single") == "下列说法正确的是（　　）",
        _g.ensure_answer_parens("下列说法正确的是（　）", "single"),
    )
    check(
        "判断括号留 2 字符",
        _g.ensure_answer_parens("同业存单属于负债业务（　）。", "judge") == "同业存单属于负债业务（　　）。",
        _g.ensure_answer_parens("同业存单属于负债业务（　）。", "judge"),
    )
    check(
        "多选括号留 6 字符",
        _g.ensure_answer_parens("下列属于流动资产的有（　）。", "multi") == "下列属于流动资产的有（　　　　　　）。",
        _g.ensure_answer_parens("下列属于流动资产的有（　）。", "multi"),
    )
    check(
        "无括号时在末尾补 2 字符括号",
        _g.ensure_answer_parens("下列说法正确的是。", "single").endswith("\u3000（　　）"),
        _g.ensure_answer_parens("下列说法正确的是。", "single"),
    )
    check(
        "无括号时在末尾补 6 字符括号",
        _g.ensure_answer_parens("下列属于流动资产的有：", "multi").endswith("\u3000（　　　　　　）"),
        _g.ensure_answer_parens("下列属于流动资产的有：", "multi"),
    )
    check(
        "括号里已有文字时不动它、另补一个空括号",
        _g.ensure_answer_parens("（2015 年修订）下列说法正确的是", "single") == "（2015 年修订）下列说法正确的是\u3000（　　）",
        _g.ensure_answer_parens("（2015 年修订）下列说法正确的是", "single"),
    )
    check("PDF 版实际渲染出 2 字符空括号", "（　　）" in html)
    check("PDF 版实际渲染出 6 字符空括号（多选）", "（　　　　　　）" in html)

    # 文件名不带模式后缀
    from services.export import exporter as _exp

    fname = _exp.default_filename(vm, "pdf")
    check(
        "默认文件名不含练习/考试字样",
        not any(w in fname for w in ("练习", "模拟", "正式考试")),
        fname,
    )

    # 四项目必须同一行、左中右三栏
    import re as _re

    m = _re.search(r'<div class="header-row">(.*?)</div>', html, _re.S)
    zones = _re.findall(r'<span class="h-(left|mid|right)">', html)
    check("卷头为一行三栏（左/中/右）", bool(m) and zones == ["left", "mid", "right"], zones)
    check(
        "左=考生信息、中=题量与满分、右=得分",
        bool(m) and "考生信息" in m.group(1).split("h-mid")[0] and "共 5 题" in m.group(1) and "得分" in m.group(1).split("h-right")[-1],
    )

    # Word 用三栏表格实现同样的对齐
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _AL

    from services.export import to_docx as _to_docx

    _p = os.path.join(SANDBOX, "check_header.docx")
    _to_docx.build_docx(vm, _p)
    import docx as _docx

    _d = _docx.Document(_p)
    _cells = _d.tables[0].rows[0].cells if _d.tables else []
    _aligns = [c.paragraphs[0].alignment for c in _cells]
    check(
        "Word 卷头三栏且左中右对齐",
        len(_cells) == 3 and _aligns == [_AL.LEFT, _AL.CENTER, _AL.RIGHT],
        [c.text for c in _cells],
    )

    # 卷头与下方横线之间不留空行
    _pf = _d.paragraphs[1].paragraph_format
    _gap = (_pf.space_before.pt if _pf.space_before else 0) + (
        getattr(_pf.line_spacing, "pt", _pf.line_spacing or 0)
    )
    check("Word 卷头紧贴下方横线（无空行）", _gap <= 2, "%.1fpt" % _gap)
    check(
        "PDF 版卷头紧贴下方横线（无空行）",
        "hr.sep { border: none; border-top: 1pt solid #000; margin: 0 0 10pt; }" in html,
    )

    # ---------------------------------------------------------------- 重难错题
    section("重难错题")
    ho = A.hard_overview({})["data"]
    check("初始没有重难错题", ho["total"] == 0, ho["total"])

    # 从正式考试会话（已交卷）里勾选两道错题
    wrong_pos = [w["pos"] for w in res["wrong_items"]]
    r1 = A.toggle_hard({"session_id": sid2, "pos": wrong_pos[0], "hard": True})
    check("勾选加入重难错题", r1["ok"] and r1["data"]["hard"] is True and r1["data"]["count"] == 1, r1.get("data"))
    r2 = A.toggle_hard({"session_id": sid2, "pos": wrong_pos[1], "hard": True})
    check("再勾一题", r2["ok"] and r2["data"]["count"] == 2, r2.get("data"))

    ho = A.hard_overview({})["data"]
    check("模块按题库汇总 1 个题库 2 道题", ho["bank_cnt"] == 1 and ho["total"] == 2, ho)
    qids = A.hard_map({})["data"]["map"].get(str(bank1), [])
    check("hard_map 返回该题库的 qid 列表", len(qids) == 2, qids)

    # 重复勾选不会重复入库
    A.toggle_hard({"session_id": sid2, "pos": wrong_pos[0], "hard": True})
    check("重复勾选不重复计数", A.hard_overview({})["data"]["total"] == 2)

    # 用重难错题出卷
    hard_settings = dict(settings, source="hard", mode="practice")
    hpv = A.preview_paper(hard_settings)
    check("重难错题可预览出卷（2 题）", hpv["ok"] and hpv["data"]["total"] == 2, hpv.get("data", {}).get("total"))
    hs = A.create_session(hard_settings)
    check("重难错题建卷成功", hs["ok"] and hs["data"]["source"] == "hard", hs.get("data", {}).get("source"))
    hsid = hs["data"]["session_id"]

    # 在练习中取消勾选 = 移出重难错题
    hitem = hs["data"]["items"][0]
    rm = A.toggle_hard({"session_id": hsid, "pos": hitem["pos"], "hard": False})
    check("练习中取消勾选即移出（剩 1 题）", rm["ok"] and rm["data"]["count"] == 1, rm.get("data"))
    A.toggle_hard({"session_id": hsid, "pos": hitem["pos"], "hard": True})
    check("可再次勾回（回到 2 题）", A.hard_overview({})["data"]["total"] == 2)

    # 练习重难错题并交卷
    with db.ro() as con:
        hrows = con.execute("SELECT pos, answer_disp_json FROM session_items WHERE session_id=? ORDER BY pos", (hsid,)).fetchall()
    for r in hrows:
        A.answer({"session_id": hsid, "pos": r["pos"], "chosen": json.loads(r["answer_disp_json"])})
    hsub = A.submit_session({"session_id": hsid})["data"]
    check("重难错题练习交卷：2 题全对", (hsub["right"], hsub["total"]) == (2, 2), (hsub["right"], hsub["total"]))

    # 重难错题三格式导出
    for fmt in ("docx", "xlsx", "pdf"):
        try:
            out = A.export_hard({"bank_id": bank1, "fmt": fmt})
            path = (out.get("data") or {}).get("path")
            size = os.path.getsize(path) if path and os.path.exists(path) else 0
            check("重难错题导出 %s" % fmt, bool(path) and size > 2000, "%s %d bytes" % (path, size))
        except Exception as e:
            check("重难错题导出 %s" % fmt, False, "%s: %s" % (type(e).__name__, e))

    # 题库删除后，重难错题一并删除
    A.delete_bank({"bank_id": bank1})
    ho2 = A.hard_overview({})["data"]
    check("删除题库后重难错题同步清空", ho2["total"] == 0 and ho2["bank_cnt"] == 0, ho2)

    # ---------------------------------------------------------------- 汇总
    print("\n================ 结果 ================")
    print("通过 %d 项，失败 %d 项" % (len(PASS), len(FAIL)))
    if FAIL:
        print("失败项：")
        for f in FAIL:
            print("  -", f)
    print("导出目录：", os.path.join(SANDBOX, "SuiTi", "exports"))
    return 0 if not FAIL else 1


if __name__ == "__main__":
    sys.exit(main())
